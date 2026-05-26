"""
Agency-ready brand deliverables — PDF (ReportLab) and PPTX (python-pptx).
Uses structured workflow output; no new cognition architecture.
"""
from __future__ import annotations

import json
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak
from reportlab.lib.enums import TA_LEFT

BRAND_NAVY = colors.HexColor("#0a1727")
BRAND_CYAN = colors.HexColor("#4fe5ff")
WHITE = colors.white

WORKFLOW_TITLES = {
    "brand_dna": "Brand DNA",
    "positioning": "Positioning",
    "positioning_engine": "Positioning Engine",
    "tone_guide": "Tone Guide",
    "messaging_pillars": "Messaging Framework",
    "messaging_framework": "Messaging Framework",
    "audience_psychology": "Audience Psychology",
    "campaign_strategy": "Campaign Direction",
    "campaign_direction": "Campaign Direction",
    "pack": "Full Brand Operating System",
    "full": "Full Brand Operating System",
}

SECTION_LABELS = {
    "brand_dna": "Brand DNA",
    "positioning": "Positioning",
    "tone_rules": "Tone Guide",
    "messaging_pillars": "Messaging Framework",
    "audience_psychology": "Audience Psychology",
    "campaign_direction": "Campaign Direction",
    "founder_narrative": "Founder Narrative",
}


def _text_from_section(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, str):
        return val.strip()
    if isinstance(val, dict):
        parts = []
        if val.get("narrative"):
            parts.append(str(val["narrative"]))
        for k, v in val.items():
            if k == "narrative":
                continue
            if isinstance(v, (str, int, float)):
                parts.append(f"{k.replace('_', ' ').title()}: {v}")
            elif isinstance(v, list):
                for item in v[:12]:
                    if isinstance(item, dict):
                        parts.append(" • ".join(f"{a}: {b}" for a, b in item.items() if b))
                    else:
                        parts.append(str(item))
        return "\n\n".join(p for p in parts if p)
    if isinstance(val, list):
        return "\n".join(_text_from_section(x) for x in val[:20])
    return str(val)


def _build_sections(package: Dict[str, Any], narrative: str = "") -> List[Tuple[str, str]]:
    sections: List[Tuple[str, str]] = []
    if narrative:
        sections.append(("Strategic Narrative", narrative))
    for key, val in (package or {}).items():
        label = SECTION_LABELS.get(key, key.replace("_", " ").title())
        body = _text_from_section(val)
        if body:
            sections.append((label, body))
    return sections


def render_brand_pdf(
    title: str,
    package: Dict[str, Any],
    narrative: str = "",
    subtitle: str = "GodFather Brand Operating System",
) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=56,
        bottomMargin=48,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BrandTitle",
        parent=styles["Heading1"],
        fontSize=22,
        textColor=BRAND_NAVY,
        spaceAfter=12,
    )
    h2_style = ParagraphStyle(
        "BrandH2",
        parent=styles["Heading2"],
        fontSize=14,
        textColor=BRAND_CYAN,
        spaceBefore=16,
        spaceAfter=8,
    )
    body_style = ParagraphStyle(
        "BrandBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        alignment=TA_LEFT,
    )

    story: List[Any] = []
    story.append(Paragraph(_escape(title), title_style))
    story.append(Paragraph(_escape(subtitle), body_style))
    story.append(Spacer(1, 0.25 * inch))

    for heading, body in _build_sections(package, narrative):
        story.append(Paragraph(_escape(heading), h2_style))
        for para in body.split("\n\n"):
            para = para.strip()
            if not para:
                continue
            story.append(Paragraph(_escape(para.replace("\n", "<br/>")), body_style))
            story.append(Spacer(1, 0.08 * inch))

    doc.build(story)
    return buffer.getvalue()


def render_brand_pptx(
    title: str,
    package: Dict[str, Any],
    narrative: str = "",
) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as e:
        raise ImportError("python-pptx required for PPT export") from e

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title[:80]
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Strategic Brand Direction — GodFather"

    sections = _build_sections(package, narrative)
    for heading, body in sections[:12]:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = heading[:60]
        tf = slide.placeholders[1].text_frame
        tf.clear()
        chunks = body[:2400].split("\n\n")
        for i, chunk in enumerate(chunks[:8]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = chunk[:500]
            p.font.size = Pt(14)

    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def _escape(s: str) -> str:
    return (
        (s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def record_export_audit(session_id: int, workflow: str, fmt: str, user=None) -> None:
    try:
        from user_sessions.services.product_signals import record_product_signal

        record_product_signal(session_id, "export", {"workflow": workflow, "format": fmt})
    except Exception:
        pass
    try:
        from user_sessions.models import BrandMemory

        row = BrandMemory.objects.filter(session_id=session_id, key="export_audit").first()
        log = list((row.value or {}).get("entries", [])) if row else []
        log.append(
            {
                "workflow": workflow,
                "format": fmt,
                "user_id": getattr(user, "id", None),
            }
        )
        log = log[-50:]
        BrandMemory.objects.update_or_create(
            session_id=session_id,
            key="export_audit",
            defaults={
                "memory_type": "brand_fact",
                "content": json.dumps(log[-1], ensure_ascii=False)[:500],
                "value": {"entries": log},
                "importance_score": 0.5,
            },
        )
    except Exception:
        pass


def workflow_display_name(workflow: str) -> str:
    return WORKFLOW_TITLES.get(workflow, workflow.replace("_", " ").title())
